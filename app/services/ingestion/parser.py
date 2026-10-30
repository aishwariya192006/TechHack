"""
BrandShield AI - Offset-Preserving AST Tokenization Engine.
Parses Raw Text, PDF (PyMuPDF), and DOCX (python-docx / OpenXML) into structured AST span maps
guaranteeing mathematical 1:1 character offset alignment (raw_text[start:end] == span.text),
bounding-box coordinates, page indices, and paragraph/table hierarchy.
"""

from dataclasses import dataclass, field
import asyncio
import hashlib
import io
import logging
import re
from typing import Any, Optional
import unicodedata
import zipfile

import pymupdf as fitz  # PyMuPDF

logger = logging.getLogger("brandshield.services.ingestion.parser")

# Ingestion Semaphore limiting concurrent CPU-bound document parsing tasks to 8
INGESTION_SEMAPHORE = asyncio.Semaphore(8)

# Sentence splitting regex matching terminal punctuation followed by whitespace or quote
SENTENCE_SPLIT_REGEX = re.compile(r"([^\.\?!]+[\.\?!]+(?:['\"\)\]\s]|\Z))", re.UNICODE)


@dataclass
class ASTSpanNode:
    """Granular AST node representing paragraph, sentence, table cell, or PDF visual block."""
    id: str
    node_type: str  # "paragraph", "sentence", "heading", "table_cell", "pdf_block"
    text: str
    start_char: int
    end_char: int
    page_number: Optional[int] = None
    bbox: Optional[list[float]] = None  # [x0, y0, x1, y1]
    style: Optional[str] = None
    paragraph_index: Optional[int] = None
    table_index: Optional[int] = None
    row_index: Optional[int] = None
    col_index: Optional[int] = None
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        """Serializes node to JSONB-compatible dictionary excluding None values."""
        data: dict[str, Any] = {
            "id": self.id,
            "node_type": self.node_type,
            "text": self.text,
            "start_char": self.start_char,
            "end_char": self.end_char,
        }
        if self.page_number is not None:
            data["page_number"] = self.page_number
        if self.bbox is not None:
            data["bbox"] = self.bbox
        if self.style is not None:
            data["style"] = self.style
        if self.paragraph_index is not None:
            data["paragraph_index"] = self.paragraph_index
        if self.table_index is not None:
            data["table_index"] = self.table_index
        if self.row_index is not None:
            data["row_index"] = self.row_index
        if self.col_index is not None:
            data["col_index"] = self.col_index
        if self.metadata:
            data["metadata"] = self.metadata
        return data


@dataclass
class ParsedDocument:
    """Container holding global normalized document text, AST span map, and parsing metrics."""
    raw_text: str
    ast_span_map: list[dict[str, Any]]
    document_type: str
    total_characters: int
    total_words: int
    total_spans: int
    metadata: dict[str, Any] = field(default_factory=dict)


class OffsetPreservingASTParser:
    """
    High-precision AST tokenizer converting multi-format assets into an offset-preserved
    span graph suitable for regulatory finding citation and evidence-linked verification.
    """

    @classmethod
    def parse_text(cls, raw_input: str, document_type: str = "TEXT") -> ParsedDocument:
        """
        Parses raw text or Markdown input:
        1. Normalizes Unicode to standard NFC form.
        2. Segments into paragraphs and nested sentences preserving exact global character offsets.
        """
        normalized_text = unicodedata.normalize("NFC", raw_input)
        spans: list[ASTSpanNode] = []

        # Split on double newlines to isolate paragraphs
        para_pattern = re.compile(r"(\S(?:.*?\S)?)(?=\n\s*\n|\Z)", re.DOTALL)

        raw_text_parts: list[str] = []
        global_offset = 0
        para_idx = 0

        # Iterate through paragraphs
        for match in para_pattern.finditer(normalized_text):
            p_text = match.group(1).strip()
            if not p_text:
                continue

            if raw_text_parts:
                raw_text_parts.append("\n\n")
                global_offset += 2

            p_start = global_offset
            raw_text_parts.append(p_text)
            p_end = global_offset + len(p_text)
            global_offset = p_end

            # Paragraph node
            p_node_type = "heading" if (p_text.startswith("#") or len(p_text) < 60 and not p_text.endswith(".")) else "paragraph"
            para_id = f"p-{para_idx}"
            spans.append(
                ASTSpanNode(
                    id=para_id,
                    node_type=p_node_type,
                    text=p_text,
                    start_char=p_start,
                    end_char=p_end,
                    paragraph_index=para_idx,
                )
            )

            # Segment sentences within paragraph
            cls._extract_sentence_spans(
                parent_text=p_text,
                parent_start=p_start,
                parent_id=para_id,
                para_index=para_idx,
                spans_dest=spans,
            )
            para_idx += 1

        full_raw_text = "".join(raw_text_parts)
        cls._verify_offset_integrity(full_raw_text, spans)

        return ParsedDocument(
            raw_text=full_raw_text,
            ast_span_map=[s.to_dict() for s in spans],
            document_type=document_type,
            total_characters=len(full_raw_text),
            total_words=len(full_raw_text.split()),
            total_spans=len(spans),
            metadata={"paragraph_count": para_idx},
        )

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> ParsedDocument:
        """
        Parses PDF using PyMuPDF (fitz):
        Extracts text block-by-block, preserving page numbers, bounding box [x0, y0, x1, y1],
        and mapping each block and sentence to global character offsets.
        """
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        spans: list[ASTSpanNode] = []
        raw_text_parts: list[str] = []
        global_offset = 0

        total_pages = len(doc)
        block_counter = 0

        for page_idx in range(total_pages):
            page = doc[page_idx]
            page_num = page_idx + 1
            # Extract blocks: (x0, y0, x1, y1, text, block_no, block_type)
            blocks = page.get_text("blocks")

            for b in blocks:
                x0, y0, x1, y1, b_text, b_no, b_type = b
                if b_type != 0:  # 0 indicates text; 1 indicates image
                    continue

                cleaned_text = unicodedata.normalize("NFC", b_text.strip())
                if not cleaned_text:
                    continue

                if raw_text_parts:
                    raw_text_parts.append("\n\n")
                    global_offset += 2

                start_char = global_offset
                raw_text_parts.append(cleaned_text)
                end_char = global_offset + len(cleaned_text)
                global_offset = end_char

                bbox = [round(x0, 2), round(y0, 2), round(x1, 2), round(y1, 2)]
                block_id = f"pdf-p{page_num}-b{b_no}"

                spans.append(
                    ASTSpanNode(
                        id=block_id,
                        node_type="pdf_block",
                        text=cleaned_text,
                        start_char=start_char,
                        end_char=end_char,
                        page_number=page_num,
                        bbox=bbox,
                        metadata={"block_number": b_no},
                    )
                )

                # Nested sentence spans with bounding box inheritance
                cls._extract_sentence_spans(
                    parent_text=cleaned_text,
                    parent_start=start_char,
                    parent_id=block_id,
                    para_index=block_counter,
                    spans_dest=spans,
                    page_number=page_num,
                    bbox=bbox,
                )
                block_counter += 1

        doc.close()
        full_raw_text = "".join(raw_text_parts)
        cls._verify_offset_integrity(full_raw_text, spans)

        return ParsedDocument(
            raw_text=full_raw_text,
            ast_span_map=[s.to_dict() for s in spans],
            document_type="PDF",
            total_characters=len(full_raw_text),
            total_words=len(full_raw_text.split()),
            total_spans=len(spans),
            metadata={"page_count": total_pages, "block_count": block_counter},
        )

    @classmethod
    def parse_docx(cls, file_bytes: bytes) -> ParsedDocument:
        """
        Parses Microsoft Word DOCX documents:
        Extracts paragraphs and table cells, capturing style names, indices,
        and global character offsets.
        """
        import importlib.util
        if importlib.util.find_spec("docx") is not None:
            try:
                return cls._parse_docx_with_library(file_bytes)
            except Exception as ex:
                logger.warning(
                    f"python-docx failed to parse document, falling back to native OpenXML: {ex}"
                )
                return cls._parse_docx_openxml(file_bytes)
        return cls._parse_docx_openxml(file_bytes)

    @classmethod
    def _parse_docx_with_library(cls, file_bytes: bytes) -> ParsedDocument:
        """Extracts paragraphs and tables using python-docx."""
        import docx  # type: ignore[import-untyped]

        doc = docx.Document(io.BytesIO(file_bytes))
        spans: list[ASTSpanNode] = []
        raw_text_parts: list[str] = []
        global_offset = 0

        p_idx = 0
        for p in doc.paragraphs:
            text = unicodedata.normalize("NFC", p.text.strip())
            if not text:
                continue

            if raw_text_parts:
                raw_text_parts.append("\n\n")
                global_offset += 2

            p_start = global_offset
            raw_text_parts.append(text)
            p_end = global_offset + len(text)
            global_offset = p_end

            style_name = p.style.name if p.style else "Normal"
            node_type = "heading" if "heading" in style_name.lower() or "title" in style_name.lower() else "paragraph"
            para_id = f"docx-p{p_idx}"

            spans.append(
                ASTSpanNode(
                    id=para_id,
                    node_type=node_type,
                    text=text,
                    start_char=p_start,
                    end_char=p_end,
                    style=style_name,
                    paragraph_index=p_idx,
                )
            )

            cls._extract_sentence_spans(
                parent_text=text,
                parent_start=p_start,
                parent_id=para_id,
                para_index=p_idx,
                spans_dest=spans,
                style=style_name,
            )
            p_idx += 1

        # Process tables
        t_idx = 0
        for table in doc.tables:
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    cell_text = unicodedata.normalize("NFC", cell.text.strip())
                    if not cell_text:
                        continue

                    if raw_text_parts:
                        raw_text_parts.append("\n\n")
                        global_offset += 2

                    c_start = global_offset
                    raw_text_parts.append(cell_text)
                    c_end = global_offset + len(cell_text)
                    global_offset = c_end

                    cell_id = f"docx-tbl{t_idx}-r{r_idx}-c{c_idx}"
                    spans.append(
                        ASTSpanNode(
                            id=cell_id,
                            node_type="table_cell",
                            text=cell_text,
                            start_char=c_start,
                            end_char=c_end,
                            table_index=t_idx,
                            row_index=r_idx,
                            col_index=c_idx,
                        )
                    )
            t_idx += 1

        full_raw_text = "".join(raw_text_parts)
        cls._verify_offset_integrity(full_raw_text, spans)

        return ParsedDocument(
            raw_text=full_raw_text,
            ast_span_map=[s.to_dict() for s in spans],
            document_type="DOCX",
            total_characters=len(full_raw_text),
            total_words=len(full_raw_text.split()),
            total_spans=len(spans),
            metadata={"paragraph_count": p_idx, "table_count": t_idx},
        )

    @classmethod
    def _parse_docx_openxml(cls, file_bytes: bytes) -> ParsedDocument:
        """Native OpenXML fallback parser using zipfile and xml.etree.ElementTree."""
        import xml.etree.ElementTree as ET

        spans: list[ASTSpanNode] = []
        raw_text_parts: list[str] = []
        global_offset = 0

        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        with zipfile.ZipFile(io.BytesIO(file_bytes), "r") as zf:
            root = ET.fromstring(zf.read("word/document.xml"))
            body = root.find("w:body", ns)
            if body is None:
                return ParsedDocument(
                    raw_text="",
                    ast_span_map=[],
                    document_type="DOCX",
                    total_characters=0,
                    total_words=0,
                    total_spans=0,
                )

            p_idx = 0
            tbl_idx = 0

            for element in body:
                tag = element.tag.split("}")[-1]

                if tag == "p":
                    # Extract paragraph text
                    text_nodes = [n.text for n in element.findall(".//w:t", ns) if n.text]
                    p_text = unicodedata.normalize("NFC", "".join(text_nodes).strip())
                    if not p_text:
                        continue

                    # Extract style
                    style_elem = element.find(".//w:pStyle", ns)
                    style = "Normal"
                    if style_elem is not None:
                        style = style_elem.attrib.get(f"{{{ns['w']}}}val", "Normal")

                    if raw_text_parts:
                        raw_text_parts.append("\n\n")
                        global_offset += 2

                    p_start = global_offset
                    raw_text_parts.append(p_text)
                    p_end = global_offset + len(p_text)
                    global_offset = p_end

                    node_type = "heading" if "heading" in style.lower() or "title" in style.lower() else "paragraph"
                    para_id = f"docx-p{p_idx}"

                    spans.append(
                        ASTSpanNode(
                            id=para_id,
                            node_type=node_type,
                            text=p_text,
                            start_char=p_start,
                            end_char=p_end,
                            style=style,
                            paragraph_index=p_idx,
                        )
                    )

                    cls._extract_sentence_spans(
                        parent_text=p_text,
                        parent_start=p_start,
                        parent_id=para_id,
                        para_index=p_idx,
                        spans_dest=spans,
                        style=style,
                    )
                    p_idx += 1

                elif tag == "tbl":
                    # Extract table rows and cells
                    for r_idx, tr in enumerate(element.findall(".//w:tr", ns)):
                        for c_idx, tc in enumerate(tr.findall(".//w:tc", ns)):
                            cell_text_nodes = [n.text for n in tc.findall(".//w:t", ns) if n.text]
                            cell_text = unicodedata.normalize("NFC", "".join(cell_text_nodes).strip())
                            if not cell_text:
                                continue

                            if raw_text_parts:
                                raw_text_parts.append("\n\n")
                                global_offset += 2

                            c_start = global_offset
                            raw_text_parts.append(cell_text)
                            c_end = global_offset + len(cell_text)
                            global_offset = c_end

                            cell_id = f"docx-tbl{tbl_idx}-r{r_idx}-c{c_idx}"
                            spans.append(
                                ASTSpanNode(
                                    id=cell_id,
                                    node_type="table_cell",
                                    text=cell_text,
                                    start_char=c_start,
                                    end_char=c_end,
                                    table_index=tbl_idx,
                                    row_index=r_idx,
                                    col_index=c_idx,
                                )
                            )
                    tbl_idx += 1

        full_raw_text = "".join(raw_text_parts)
        cls._verify_offset_integrity(full_raw_text, spans)

        return ParsedDocument(
            raw_text=full_raw_text,
            ast_span_map=[s.to_dict() for s in spans],
            document_type="DOCX",
            total_characters=len(full_raw_text),
            total_words=len(full_raw_text.split()),
            total_spans=len(spans),
            metadata={"paragraph_count": p_idx, "table_count": tbl_idx},
        )

    @classmethod
    def _extract_sentence_spans(
        cls,
        parent_text: str,
        parent_start: int,
        parent_id: str,
        para_index: int,
        spans_dest: list[ASTSpanNode],
        page_number: Optional[int] = None,
        bbox: Optional[list[float]] = None,
        style: Optional[str] = None,
    ) -> None:
        """
        Segments sentence units within a parent text block while computing exact character offsets.
        """
        sentence_matches = list(re.finditer(r"([^\.\?!]+[\.\?!]+|.+)", parent_text))
        s_idx = 0

        for match in sentence_matches:
            s_raw = match.group(0).strip()
            if not s_raw:
                continue

            # Compute relative offsets
            rel_start = match.start()
            # Account for stripped whitespace at beginning of sentence
            rel_start += match.group(0).find(s_raw)
            rel_end = rel_start + len(s_raw)

            abs_start = parent_start + rel_start
            abs_end = parent_start + rel_end

            spans_dest.append(
                ASTSpanNode(
                    id=f"{parent_id}-s{s_idx}",
                    node_type="sentence",
                    text=s_raw,
                    start_char=abs_start,
                    end_char=abs_end,
                    page_number=page_number,
                    bbox=bbox,
                    style=style,
                    paragraph_index=para_index,
                    metadata={"sentence_index": s_idx},
                )
            )
            s_idx += 1

    @staticmethod
    def _verify_offset_integrity(raw_text: str, spans: list[ASTSpanNode]) -> None:
        """
        Mathematical invariant check:
        Asserts that every node's declared text strictly equals raw_text[start:end].
        """
        for span in spans:
            slice_text = raw_text[span.start_char : span.end_char]
            if slice_text != span.text:
                raise ValueError(
                    f"AST span offset alignment fault in node '{span.id}': "
                    f"expected '{span.text[:30]}...', got '{slice_text[:30]}...'"
                )


# ------------------------------------------------------------------------------
# Non-blocking CPU offloading wrappers (asyncio.to_thread + backpressure semaphore)
# ------------------------------------------------------------------------------

async def parse_pdf_async(file_bytes: bytes) -> ParsedDocument:
    """Non-blocking PDF parser executing fitz extraction in thread pool under Semaphore(8)."""
    async with INGESTION_SEMAPHORE:
        return await asyncio.to_thread(OffsetPreservingASTParser.parse_pdf, file_bytes)


async def parse_docx_async(file_bytes: bytes) -> ParsedDocument:
    """Non-blocking DOCX parser executing OpenXML extraction in thread pool under Semaphore(8)."""
    async with INGESTION_SEMAPHORE:
        return await asyncio.to_thread(OffsetPreservingASTParser.parse_docx, file_bytes)


async def parse_text_async(raw_input: str, document_type: str = "TEXT") -> ParsedDocument:
    """Non-blocking plain text parser executing AST tokenization in thread pool under Semaphore(8)."""
    async with INGESTION_SEMAPHORE:
        return await asyncio.to_thread(OffsetPreservingASTParser.parse_text, raw_input, document_type)


async def compute_sha256_async(data: bytes) -> str:
    """Non-blocking SHA-256 fingerprint computation offloaded to thread pool."""
    return await asyncio.to_thread(lambda: hashlib.sha256(data).hexdigest())
